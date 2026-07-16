import warnings
import pandas as pd
from openpyxl import load_workbook
from pandas.core.interchange import column

warnings.filterwarnings(
    "ignore",
    message="Print area cannot be set*"
)

warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported*"
)

df = pd.read_excel('./src/catalogueOutils.xlsx')
df = df.drop(columns=['type de contenu', 'étape parcours','Socle', 'Créé dans Up2C',
       'A jour dans Up2C', 'Besoin RH', 'Besoin ordi', 'besoin matériel','A jour dans Sharepoint',
       'compétences', 'Unnamed: 17'])
df = df.rename(columns={'Nom de l\'outil':'name', 'programmation':'planning', 'up2c-projet':'project', 'up2c-séquence':'sequence',
       'modalité':'mode', 'objectifs généraux':'target', 'chemin':'path'})

wb = load_workbook('./src/catalogueOutils.xlsx')
ws = wb ["catalogue"]

urls = []

for cell in ws["O"][1:]:
    urls.append(cell.hyperlink.target if cell.hyperlink else None)

df["path"] = urls


apc_sheets = [
    "Socle (APC)",
    "Boite à outils (APC)",
    "Intégration SAS (APC)",
    "Parcours Métier (APC)"
]

skills_map = {}

for sheet in apc_sheets:

    apc = pd.read_excel(
        "./src/catalogueOutils.xlsx",
        sheet_name=sheet
    )

    skill_cols = [
        col
        for col in apc.columns
        if isinstance(col, str)
        and len(col) == 2
        and col[0].isdigit()
    ]

    for _, row in apc.iterrows():

        activity = row["Séances"]

        if pd.isna(activity):
            continue

        skills = [
            skill
            for skill in skill_cols
            if str(row[skill]).strip().lower() == "oui"
        ]

        skills_map[str(activity).strip()] = skills


df["skills"] = (
    df["name"]
    .astype(str)
    .str.strip()
    .map(skills_map)
)



df["skills"] = df["skills"].apply(
    lambda x: x if isinstance(x, list) else []
)

#df.to_csv('./output/tools.csv', index=False)

df_stud = pd.read_excel('./src/stagiaires.xlsx')

df_stud = df_stud.drop(columns=['ID', 'Civilité', 'Téléphone', 'Type de contrats', 'Période',
       'Groupe', 'Parcours', 'Projet', 'Âge',
       'Âge à l\'entrée', 'Date de Naissance', 'Nom de Naissance',
       'Lieu de Naissance', 'Nationalité', 'Document d\'identité valide',
       'Adresse', 'Code Postal', 'Ville', 'QPV', 'Code QPV', 'ZRR',
       'Contrat CEJ', 'Niveau de Diplôme', 'Expérience Professionnelle', 'CIR',
       'RQTH', 'Demandeur d\'emploi longue durée', 'Prise en charge ASE',
       'Sous main de justice', 'Minutes en Centre', 'Minutes en Stage',
       'Minutes d\'Absences',
       'Date de signature du contrat', 'Date de sortie prévisionnelle',
       'Date de sortie de formation', 'Date de départ stagiaire',
       'Date de départ école', 'Type de sortie', 'Date de première entretien',
       'Type de prescripteur', 'Contact prescripteur', 'Date de prescription',
       'Email du prescripteur', 'Téléphone du prescripteur',
       'Responsable légal - Valide', 'Responsable légal - Nom',
       'Responsable légal - Prénom', 'Responsable légal - Adresse',
       'Responsable légal - Téléphone', 'Responsable légal - Email',
       'Code Région (SoWeSign)', 'Date d\'inscription France Travail',
       'Numéro France Travail', 'Azure ID', 'Statut invitation Azure',
       'Projet professionnel final', 'mail alternatif'])
df_stud = df_stud.rename(columns={'Nom':'lastname', 'Prénom':'firstname', 'Email':'email', 'École':'structure', 'Site':'site',
       'Projets professionnels':'project', 'Formateur Référent':'referent',
       'Date d\'entrée en formation':'starting_date', 'lien One Drive':'drive_path'})


df_stud.loc[
    df_stud["referent"] == "JAMIL HUSSAIN Sura",
    "referent"
] = "JAMIL-HUSSAIN Sura"

df_stud.loc[
    df_stud["referent"] == "LEON QUERO German",
    "referent"
] = "LEON-QUERO German"


#df_stud.to_csv('./output/students.csv', index=False)

df_staff = pd.read_csv('./src/users.csv')
df_staff = df_staff.drop(columns=['Display name', 'DirSyncEnabled', 'Object Id','When created', 'Soft deletion time stamp',
       'AgeGroup', 'Department', 'Preferred data location',
       'CountryOrRegion', 'Office', 'StateOrProvince', 'Usage location',
       'Last dirsync time', 'Block credential', 'Licenses',
       'Password never expires', 'Last password change time stamp',
       'Mobile Phone', 'Phone number', 'Postal code', 'Preferred language',
       'Street address', 'Fax', 'Proxy addresses'])
df_staff = df_staff.rename(columns={'User principal name' : 'email', 'First name':'firstname', 'Last name':'lastname',
       'City' :'site', 'Title':'role'})
df_staff["full_name"]=(df_staff['lastname'].str.strip().str.upper()+ " " + df_staff['firstname'])

referents = (
       df_stud[["site", "referent"]].drop_duplicates(subset=['referent'])
)

df_staff = df_staff.merge(referents, left_on='full_name', right_on='referent', how='left')

mask = df_staff["referent"].notna()

df_staff.loc[mask, "role"]="USER_REF"
df_staff.loc[mask, "site"] = df_staff.loc[mask, "site_y"]

df_staff = df_staff.drop(columns=['referent', 'site_y', 'site_x'])

coords = {
       "JACOB Thibault" : "Roubaix",
       "MATTON Julie" : "Saint Omer",
       "MILLET Lucie" : "Armentières",
       "MAUNOURY Franklin" : "Lille",
       "PICART Agnès" : "Siège",
       "CATHELAIN Amélie": "Siège",
       "LAMBET Christine": "Siège",
       "DEWITTE Eglantine": "Siège",
       "BEN ZARB Ezzeddine": "Roubaix",
       "DECROIX Jean-Pierre": "Siège",
       "BOLPAIRE Jérôme" : "Armentières",
       "GRIS Laurence" : "Lille",
       "BURCKER Olivier" : "Siège",
       "CAROULLE Xavier" : "Saint Omer"
}

for user, locus in coords.items():
       df_staff.loc[df_staff["full_name"] == user, ["site", "role"]] = [locus, "USER_COORD"]

df_staff.loc[df_staff['full_name']=='CARON Delphine', ['role', 'site']]=['USER_REF', 'Saint Omer']

df_staff.loc[df_staff['full_name']=='DE GORI Jeanne', ['role', 'site']]=['USER_REF', 'Lille']

roles = ('USER_REF', 'USER_COORD')

df_staff = df_staff[df_staff['role'].isin(roles)]

#df_staff.to_csv('./output/staff.csv', index=False)

df.to_json(
    './output/tools.json',
    orient='records',
    force_ascii=False,
    indent=2
)

df_stud.to_json(
    './output/students.json',
    orient='records',
    force_ascii=False,
    indent=2,
    date_format='iso'
)

df_staff.to_json(
    './output/staff.json',
    orient='records',
    force_ascii=False,
    indent=2
)
